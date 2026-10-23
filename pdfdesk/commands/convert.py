"""Convert tab (plan phase 10): create PDFs from other files, export PDFs to other formats."""
from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
from pathlib import Path

import pymupdf

from .pages import parse_range
from .registry import command

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".gif", ".webp", ".jxr", ".pnm", ".jp2"}
OFFICE_EXT = {".doc", ".docx", ".odt", ".rtf", ".xls", ".xlsx", ".ods", ".csv", ".ppt", ".pptx", ".odp",
              ".vsd", ".vsdx", ".odg", ".wpd", ".pages", ".key", ".numbers", ".wps", ".et", ".dps"}
TEXT_EXT = {".txt", ".md", ".log", ".py", ".json", ".xml", ".ini", ".cfg"}
HTML_EXT = {".html", ".htm", ".xhtml"}
SUPPORTED = IMAGE_EXT | OFFICE_EXT | TEXT_EXT | HTML_EXT | {".pdf", ".xps", ".epub", ".cbz", ".fb2", ".svg"}


def _soffice() -> str:
    for c in ("soffice", "libreoffice"):
        p = shutil.which(c)
        if p:
            return p
    raise RuntimeError("LibreOffice is not installed (needed for Office files)")


def office_to_pdf(path: str, outdir: str, timeout: int = 240) -> str:
    profile = Path.home() / ".cache" / "pdfdesk" / "lo-profile"
    profile.mkdir(parents=True, exist_ok=True)
    cmd = [_soffice(), f"-env:UserInstallation=file://{profile}", "--headless", "--norestore",
           "--convert-to", "pdf", "--outdir", outdir, path]
    res = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
    out = Path(outdir) / (Path(path).stem + ".pdf")
    if not out.exists():
        raise RuntimeError(f"LibreOffice could not convert {path}: {res.stderr.strip() or res.stdout.strip()}")
    return str(out)


@command("Convert.FromFile")
def from_file(path: str, out: str | None = None) -> pymupdf.Document:
    """Create a PDF from any supported file. Returns the document (and saves it if out is given)."""
    ext = Path(path).suffix.lower()
    if ext not in SUPPORTED:
        raise ValueError(f"unsupported file type: {ext}")
    if ext == ".pdf":
        doc = pymupdf.open(path)
    elif ext in OFFICE_EXT:
        with tempfile.TemporaryDirectory() as td:
            with open(office_to_pdf(path, td), "rb") as f:
                doc = pymupdf.open("pdf", f.read())
    elif ext in TEXT_EXT:
        doc = _text_to_pdf(Path(path).read_text(errors="replace"))
    elif ext in HTML_EXT:
        doc = _html_to_pdf(Path(path).read_text(errors="replace"), os.path.dirname(os.path.abspath(path)))
    else:  # images and other MuPDF-readable formats (multi-frame TIFF included)
        doc = pymupdf.open("pdf", pymupdf.open(path).convert_to_pdf())
    if out:
        doc.save(out, garbage=3, deflate=True)
    return doc


def _text_to_pdf(text: str, fontsize: float = 10) -> pymupdf.Document:
    from ..core import fonts
    doc = pymupdf.open()
    ff = fonts.find("DejaVu Sans Mono") or fonts.unicode_font()
    lines = text.expandtabs(4).splitlines() or [""]
    per_page = int((842 - 100) / (fontsize * 1.3))
    for s in range(0, len(lines), per_page):
        p = doc.new_page(width=595, height=842)
        p.insert_font(fontname="mono", fontfile=ff)
        y = 50 + fontsize
        for ln in lines[s:s + per_page]:
            p.insert_text((50, y), ln[:110], fontname="mono", fontsize=fontsize)
            y += fontsize * 1.3
    return doc


def _html_to_pdf(html: str, base: str) -> pymupdf.Document:
    story = pymupdf.Story(html=html, archive=pymupdf.Archive(base))
    buf = __import__("io").BytesIO()
    writer = pymupdf.DocumentWriter(buf)
    mediabox = pymupdf.paper_rect("a4")
    where = mediabox + (40, 40, -40, -40)
    more = True
    while more:
        dev = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
    writer.close()
    return pymupdf.open("pdf", buf.getvalue())


@command("Convert.Combine")
def combine(paths: list[str], out: str | None = None) -> pymupdf.Document:
    """Combine files of any supported type into one PDF, one top-level bookmark per file."""
    result = pymupdf.open()
    toc = []
    for p in paths:
        src = from_file(p)
        start = result.page_count
        result.insert_pdf(src)
        toc.append([1, Path(p).stem, start + 1])
        toc += [[lvl + 1, t, pg + start] for lvl, t, pg, *_ in src.get_toc()]
    result.set_toc(toc)
    if out:
        result.save(out, garbage=3, deflate=True)
    return result


# ---- export -----------------------------------------------------------------------------------------------

@command("Convert.ToImages")
def to_images(doc, outdir: str, fmt: str = "png", dpi: int = 150, pages: str | None = None) -> list[str]:
    os.makedirs(outdir, exist_ok=True)
    stem = Path(doc.name).stem if doc.name else "page"
    out = []
    for i in parse_range(pages, doc.page_count):
        pix = doc[i].get_pixmap(dpi=dpi, alpha=False)
        path = os.path.join(outdir, f"{stem}_{i + 1:03d}.{fmt}")
        if fmt in ("jpg", "jpeg"):
            pix.save(path, jpg_quality=90)
        elif fmt in ("tif", "tiff"):
            from PIL import Image
            Image.frombytes("RGB", (pix.width, pix.height), pix.samples).save(path, dpi=(dpi, dpi))
        else:
            pix.save(path)
        out.append(path)
    return out


@command("Convert.ToText")
def to_text(doc, out: str, pages: str | None = None):
    with open(out, "w", encoding="utf-8") as f:
        for i in parse_range(pages, doc.page_count):
            f.write(doc[i].get_text(sort=True))
            f.write("\f")
    return out


@command("Convert.ToHtml")
def to_html(doc, out: str, pages: str | None = None):
    parts = [doc[i].get_text("html") for i in parse_range(pages, doc.page_count)]
    with open(out, "w", encoding="utf-8") as f:
        f.write("<!DOCTYPE html><html><head><meta charset='utf-8'></head><body>\n" + "\n".join(parts) + "</body></html>")
    return out


@command("Convert.ToWord")
def to_word(src: str, out: str, pages: str | None = None):
    from pdf2docx import Converter
    d = pymupdf.open(src)
    idx = parse_range(pages, d.page_count)
    cv = Converter(src)
    try:
        cv.convert(out, pages=idx)
    finally:
        cv.close()
    return out


@command("Convert.ToRtf")
def to_rtf(src: str, out: str):
    with tempfile.TemporaryDirectory() as td:
        docx = os.path.join(td, Path(out).stem + ".docx")
        to_word(src, docx)
        profile = Path.home() / ".cache" / "pdfdesk" / "lo-profile"
        subprocess.run([_soffice(), f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "rtf",
                        "--outdir", td, docx], capture_output=True, timeout=240)
        shutil.move(os.path.join(td, Path(out).stem + ".rtf"), out)
    return out


@command("Convert.ToExcel")
def to_excel(doc, out: str, pages: str | None = None):
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    n = 0
    for i in parse_range(pages, doc.page_count):
        for k, t in enumerate(doc[i].find_tables().tables):
            ws = wb.create_sheet(f"Page{i + 1}-Table{k + 1}"[:31])
            for row in t.extract():
                ws.append([("" if c is None else c) for c in row])
            n += 1
    if n == 0:  # no tables: one sheet per page, one row per text line
        for i in parse_range(pages, doc.page_count):
            ws = wb.create_sheet(f"Page{i + 1}")
            for ln in doc[i].get_text(sort=True).splitlines():
                ws.append([ln])
    wb.save(out)
    return n


@command("Convert.ToPowerPoint")
def to_powerpoint(doc, out: str, pages: str | None = None):
    """One slide per page: the page without its text as the picture, text blocks as editable text boxes."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Pt
    prs = Presentation()
    first = doc[0].rect
    emu = 12700  # EMU per point
    prs.slide_width, prs.slide_height = Emu(int(first.width * emu)), Emu(int(first.height * emu))
    blank = prs.slide_layouts[6]
    tmp = pymupdf.open()
    tmp.insert_pdf(doc)
    with tempfile.TemporaryDirectory() as td:
        for i in parse_range(pages, doc.page_count):
            page = tmp[i]
            blocks = doc[i].get_text("dict", sort=True)["blocks"]
            page.add_redact_annot(page.rect, fill=False, cross_out=False)
            page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE,
                                  graphics=pymupdf.PDF_REDACT_LINE_ART_NONE)
            png = os.path.join(td, f"{i}.png")
            page.get_pixmap(dpi=150, alpha=False).save(png)
            s = prs.slides.add_slide(blank)
            s.shapes.add_picture(png, 0, 0, prs.slide_width, prs.slide_height)
            for b in blocks:
                if b.get("type") != 0:
                    continue
                x0, y0, x1, y1 = b["bbox"]
                tb = s.shapes.add_textbox(Emu(int(x0 * emu)), Emu(int(y0 * emu)), Emu(int((x1 - x0) * emu) + emu * 4),
                                          Emu(int((y1 - y0) * emu)))
                tf = tb.text_frame
                tf.word_wrap = True
                for li, ln in enumerate(b["lines"]):
                    para = tf.paragraphs[0] if li == 0 else tf.add_paragraph()
                    for sp in ln["spans"]:
                        run = para.add_run()
                        run.text = sp["text"]
                        run.font.size = Pt(max(1, round(sp["size"])))
                        c = sp["color"]
                        run.font.color.rgb = RGBColor(c >> 16 & 255, c >> 8 & 255, c & 255)
                        run.font.bold = bool(sp["flags"] & 16)
                        run.font.italic = bool(sp["flags"] & 2)
                for m in (tf,):
                    m.margin_left = m.margin_right = m.margin_top = m.margin_bottom = 0
    prs.save(out)
    return out


@command("Convert.ToPDFA")
def to_pdfa(src: str, out: str, part: int = 2):
    """Convert to PDF/A-1b/2b/3b with OCRmyPDF's Ghostscript pipeline (no OCR on pages that have text)."""
    import ocrmypdf
    ocrmypdf.ocr(src, out, output_type=f"pdfa-{part}", skip_text=True, optimize=0, progress_bar=False,
                 tesseract_timeout=0)
    return out
