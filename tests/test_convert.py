import shutil

import pymupdf
import pytest

from pdfdesk.commands import run


def test_from_image_text_html(tmp_path, v1):
    png = tmp_path / "a.png"
    v1[0].get_pixmap(dpi=72).save(png)
    assert run("Convert.FromFile", str(png)).page_count == 1
    t = tmp_path / "a.txt"
    t.write_text("line one\nline two\n")
    assert "line two" in run("Convert.FromFile", str(t))[0].get_text()
    h = tmp_path / "a.html"
    h.write_text("<h1>Title</h1><p>Body paragraph</p>")
    assert "Body paragraph" in run("Convert.FromFile", str(h))[0].get_text()


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice not installed")
def test_from_office_and_combine(tmp_path, v1_path):
    from docx import Document
    p = tmp_path / "letter.docx"
    doc = Document()
    doc.add_heading("Letter heading", 1)
    doc.add_paragraph("Dear reader, this came from Word.")
    doc.save(p)
    pdf = run("Convert.FromFile", str(p))
    assert "this came from Word" in pdf[0].get_text()
    c = run("Convert.Combine", [v1_path, str(p)])
    assert c.page_count == 7 and [t[1] for t in c.get_toc() if t[0] == 1][:2] == ["v1", "letter"]


def test_export_formats(tmp_path, v1, v1_path):
    imgs = run("Convert.ToImages", v1, str(tmp_path / "img"), dpi=72, pages="1-2")
    assert len(imgs) == 2 and pymupdf.open(imgs[0])[0].rect.width == 595
    run("Convert.ToText", v1, str(tmp_path / "a.txt"))
    assert "Acme Ltd" in (tmp_path / "a.txt").read_text()
    run("Convert.ToHtml", v1, str(tmp_path / "a.html"))
    assert "Acme Ltd" in (tmp_path / "a.html").read_text()
    run("Convert.ToWord", v1_path, str(tmp_path / "a.docx"))
    from docx import Document
    assert "Acme Ltd" in "\n".join(p.text for p in Document(str(tmp_path / "a.docx")).paragraphs)
    run("Convert.ToExcel", v1, str(tmp_path / "a.xlsx"))
    from openpyxl import load_workbook
    assert load_workbook(tmp_path / "a.xlsx").sheetnames
    run("Convert.ToPowerPoint", v1, str(tmp_path / "a.pptx"))
    from pptx import Presentation
    prs = Presentation(str(tmp_path / "a.pptx"))
    assert len(prs.slides) == 6
    assert "Acme Ltd" in " ".join(s.text_frame.text for s in prs.slides[0].shapes if s.has_text_frame)


def test_pdfa(tmp_path, v1_path):
    out = str(tmp_path / "a.pdf")
    run("Convert.ToPDFA", v1_path, out)
    xmp = pymupdf.open(out).get_xml_metadata()
    assert "pdfaid:part" in xmp and "Acme Ltd" in pymupdf.open(out)[0].get_text()
